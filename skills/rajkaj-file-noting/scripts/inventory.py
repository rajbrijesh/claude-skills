#!/usr/bin/env python3
"""
Inventory the files in a पत्रावली before reading them.

For each file: type, page count, and (for PDFs) how each page should be read:
  text    - usable Unicode text layer (pdftotext is enough)
  english - text layer that is ordinary English
  legacy  - Latin gibberish, probably a legacy Hindi font (Kruti Dev / DevLys /
            Chanakya); treat as an image and transcribe or read visually
  scan    - little or no text layer; needs the transcribe skill or a visual read

Usage:
    python inventory.py file1.pdf photo.jpg letter.docx sheet.xlsx ...
    python inventory.py some_folder/
"""

import re
import subprocess
import sys
from pathlib import Path

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".heic", ".avif"}
ENGLISH_WORDS = {
    "the", "of", "and", "to", "in", "is", "for", "with", "on", "by", "be", "this",
    "that", "are", "as", "at", "from", "dated", "letter", "government", "department",
    "sub", "subject", "ref", "reference", "office", "no", "your", "has", "have", "will",
    "please", "sir", "under", "an", "or", "it", "not", "which", "may", "shall",
}
# Very common Hindi words as they come out of Kruti Dev / DevLys text layers
# (का के की में है और से को कि यह इस पर भी दिनांक क्रमांक विषय राज्य सरकार कार्य विभाग)
LEGACY_WORDS = {
    "dk", "ds", "dh", "esa", "gs", "gS", "vkSj", "ls", "dks", "fd", ";g", "bl",
    "ij", "Hkh", "fnukad", "Øekad", "fo\"k;", "jkT;", "ljdkj", "dk;Z", "foHkkx",
    "jktLFkku", "xzke", "iapk;r", "vf/kdkjh", ",oa", "gSA", "x;k", "fy;s", "fy,",
}
DEVANAGARI = re.compile(r"[ऀ-ॿ]")
LATIN = re.compile(r"[A-Za-z]")
MIN_CHARS = 40  # below this a page is treated as having no real text layer


def run(cmd):
    try:
        return subprocess.run(cmd, capture_output=True, text=True, timeout=120).stdout
    except Exception:
        return ""


def classify_page(text):
    stripped = re.sub(r"\s+", "", text)
    if len(stripped) < MIN_CHARS:
        return "scan"
    dev = len(DEVANAGARI.findall(stripped))
    lat = len(LATIN.findall(stripped))
    if dev >= 0.3 * len(stripped) or dev > lat:
        return "text"
    raw = text.split()
    words = re.findall(r"[A-Za-z]+", text.lower())
    if not raw or not words:
        return "scan"
    legacy_hits = sum(1 for w in raw if w.strip(",.-") in LEGACY_WORDS)
    if legacy_hits >= 3 or legacy_hits / len(raw) >= 0.05:
        # Even when some English is mixed in, Hindi body text in a legacy font
        # must be read visually; its text layer cannot be trusted.
        return "legacy"
    hits = sum(1 for w in words if w in ENGLISH_WORDS)
    # Real English prose has many common words; legacy-font Hindi has almost none.
    if hits / len(words) >= 0.08:
        return "english"
    return "legacy"


def pdf_info(path):
    info = run(["pdfinfo", str(path)])
    m = re.search(r"^Pages:\s+(\d+)", info, re.M)
    pages = int(m.group(1)) if m else 0
    kinds = []
    for i in range(1, pages + 1):
        txt = run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), str(path), "-"])
        kinds.append(classify_page(txt))
    return pages, kinds


def compress(kinds):
    """[text,text,scan,scan,scan] -> 'text: 1-2; scan: 3-5'"""
    if not kinds:
        return ""
    groups, start = [], 0
    for i in range(1, len(kinds) + 1):
        if i == len(kinds) or kinds[i] != kinds[start]:
            rng = f"{start + 1}" if i - 1 == start else f"{start + 1}-{i}"
            groups.append((kinds[start], rng))
            start = i
    merged = {}
    for kind, rng in groups:
        merged.setdefault(kind, []).append(rng)
    return "; ".join(f"{k}: {', '.join(v)}" for k, v in merged.items())


def describe(path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        pages, kinds = pdf_info(path)
        if pages == 0:
            return "pdf", "?", "could not read (encrypted or damaged?)"
        needs = [k for k in kinds if k in ("scan", "legacy")]
        note = compress(kinds)
        if needs:
            note += f"  -> {len(needs)} page(s) need OCR or a visual read"
        return "pdf", pages, note
    if ext in IMAGE_EXT:
        return "image", 1, "needs OCR or a visual read"
    if ext == ".docx":
        try:
            import docx  # python-docx
            d = docx.Document(str(path))
            paras = sum(1 for p in d.paragraphs if p.text.strip())
            body = " ".join(p.text for p in d.paragraphs)
            kind = classify_page(body) if body.strip() else "empty"
            return "word", "-", f"{paras} paragraphs, {len(d.tables)} tables, text: {kind}"
        except Exception as e:
            return "word", "-", f"read with docx skill ({e.__class__.__name__})"
    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
        try:
            import openpyxl
            if ext == ".csv":
                return "table", "-", "csv"
            wb = openpyxl.load_workbook(str(path), read_only=True)
            sheets = ", ".join(f"{ws.title} ({ws.max_row}x{ws.max_column})" for ws in wb.worksheets)
            return "excel", "-", sheets
        except Exception as e:
            return "excel", "-", f"read with xlsx skill ({e.__class__.__name__})"
    if ext == ".doc":
        return "word (old .doc)", "-", "convert with libreoffice --headless --convert-to docx"
    return ext.lstrip(".") or "unknown", "-", "check manually"


def collect(args):
    files = []
    for a in args:
        p = Path(a)
        if p.is_dir():
            files.extend(sorted(x for x in p.iterdir() if x.is_file()))
        elif p.exists():
            files.append(p)
        else:
            print(f"NOT FOUND: {a}")
    return files


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    files = collect(sys.argv[1:])
    total = 0
    print(f"{'#':>2}  {'file':40} {'type':8} {'pages':>5}  reading notes")
    for n, f in enumerate(files, 1):
        kind, pages, note = describe(f)
        if isinstance(pages, int):
            total += pages
        print(f"{n:>2}  {f.name[:40]:40} {kind:8} {str(pages):>5}  {note}")
    print(f"\n{len(files)} file(s), {total} PDF/image page(s) counted.")
    print("Page kinds: text = Unicode text layer, english = English text layer,")
    print("legacy = legacy Hindi font gibberish (treat as image), scan = no text layer.")


if __name__ == "__main__":
    main()
